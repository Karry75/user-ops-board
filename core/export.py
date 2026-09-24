"""Excel 导出（openpyxl）：押金预警 / 失败归因 / 沉默低频线索 / 用户 360。"""
from __future__ import annotations

import os
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import settings
from core import common, silent, user360
from core.common import Filters

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(color="FFFFFF", bold=True)


def _write(ws, headers: List[str], rows: List[List[Any]], widths: List[int] | None = None) -> None:
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths[i - 1] if widths and i <= len(widths) else 16)
    ws.freeze_panes = "A2"


def _out_path(name: str) -> str:
    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "output")
    os.makedirs(out_dir, exist_ok=True)
    return os.path.join(out_dir, name)


def export(kind: str, f: Filters, keyword: str = "", drill: str = "") -> str:
    wb = Workbook()
    ws = wb.active
    stamp = common.ms_to_date(common.now_ms()).replace("-", "")
    if kind == "deposit_warn":
        from core import deposit
        rows = deposit.warn_list(f, size=200000, cap=200000, drill=drill)["rows"]
        ws.title = "押金划扣预警明细"
        _write(ws, ["用户协议ID", "用户ID", "用户姓名", "用户签约手机号", "用户当前手机号", "城市",
                    "电池产品", "押金金额(元)", "押金订单创建时间", "押金创建时间", "押金到期时间",
                    "押金预警", "划扣状态", "押金划扣失败原因", "失败证据", "建议划扣金额(元)",
                    "电池SN", "电池定位地址", "电池状态", "电池在线状态", "套餐情况",
                    "协议状态", "预警层级", "最近接待备注"],
               [[r.get("id"), r.get("user_id"), r.get("user_name"), r.get("user_phone"),
                 r.get("current_phone"), r.get("sys_city_name"),
                 r.get("product_name"), r.get("deposit_fee_yuan"),
                 r.get("deposit_order_create_str"), r.get("deposit_bind_str"),
                 r.get("deposit_expire_str"), r.get("warn_text"),
                 r.get("deduct_status_cn") or r.get("deduct_status"),
                 r.get("failure_name"), r.get("failure_evidence"), r.get("suggest_fee_yuan"),
                 r.get("battery_sn"), r.get("battery_location"), r.get("battery_status_cn"),
                 r.get("battery_online_cn"), r.get("package_text"),
                 r.get("agreement_status_cn") or r.get("status"), r.get("stage_cn"),
                 r.get("note")] for r in rows],
               [12, 12, 12, 16, 16, 14, 16, 12, 20, 20, 20, 26, 12, 18, 32, 16,
                22, 34, 12, 12, 40, 12, 10, 40])
        name = f"押金划扣预警明细_{stamp}.xlsx"
    elif kind == "deposit_suggest":
        from core import deposit
        rows = deposit.suggest_list(f, size=200000, cap=200000)["rows"]
        ws.title = "失败归因与建议划扣"
        _write(ws, ["协议ID", "用户ID", "用户姓名", "手机号", "城市", "电池产品", "押金(元)", "剩余天数",
                    "失败归因", "近90天消费(元)", "近90天换电次数", "评估等级", "可划扣比例",
                    "建议划扣(元)", "建议文案"],
               [[r.get("id"), r.get("user_id"), r.get("user_name"), r.get("user_phone"),
                 r.get("sys_city_name"), r.get("product_name"), r.get("deposit_fee_yuan"),
                 r.get("days_left"), r.get("failure_name"),
                 round(int(r.get("consume_90d_fee") or 0) / 100.0, 2), r.get("exchange_90d"),
                 r.get("level_name"), r.get("suggest_ratio"), r.get("suggest_fee_yuan"),
                 r.get("suggest_text")] for r in rows],
               [12, 12, 12, 14, 14, 16, 10, 10, 16, 14, 12, 18, 10, 12, 40])
        name = f"押金失败归因与建议划扣_{stamp}.xlsx"
    elif kind in ("silent", "lowfreq"):
        from core import silent as silent_mod
        if kind == "silent":
            rows = silent_mod.silent_list(f, size=200000, cap=200000, keyword=keyword)["rows"]
            title = "沉默用户与电池线索"
        else:
            rows = silent_mod.low_freq_list(f, size=200000, cap=200000, keyword=keyword)["rows"]
            title = "低频用户与电池线索"
        ws.title = title
        silent_mod._enrich_page(rows, f)
        _write(ws, ["用户协议ID", "用户ID", "姓名", "用户签约手机号", "用户当前手机号", "城市",
                    "电池产品", "层级", "协议状态",
                    "生效天数", "租期到期", "欠租", "押金到期时间", "押金预警", "划扣状态",
                    "押金划扣失败原因", "套餐情况",
                    "持有电池SN", "电池在线", "电量(%)", "最后定位地址",
                    "最后换电网点", "最后换电时间", "是否已归还柜机", "近15/30/60/90天换电",
                    "月均频次", "动作", "动作原因"],
               [[r.get("id"), r.get("user_id"), r.get("user_name"), r.get("user_phone"),
                 r.get("current_phone"), r.get("sys_city_name"), r.get("product_name"),
                 r.get("level") or "沉默",
                 r.get("agreement_status_cn") or r.get("status"), r.get("active_days"),
                 r.get("rent_expire_str"),
                 "是" if r.get("rent_overdue") else "否", r.get("deposit_expire_str"),
                 r.get("deposit_warn_text"), r.get("deduct_status_cn"),
                 r.get("failure_name"), r.get("package_text"),
                 r.get("battery_sn"),
                 r.get("online_status_cn") or r.get("online_status"), r.get("power"),
                 r.get("last_location_address"),
                 r.get("last_site"), r.get("last_take_time_str"),
                 "是" if r.get("last_back_exchange_sn") else "否",
                 f"{r.get('n15',0)}/{r.get('n30',0)}/{r.get('n60',0)}/{r.get('n90',0)}",
                 r.get("monthly_freq"), r.get("action") or "", r.get("action_reason") or ""] for r in rows],
               [12, 12, 12, 16, 16, 14, 16, 8, 10, 10, 18, 8, 20, 26, 12, 18, 40,
                22, 10, 8, 34, 18, 18, 14, 18, 10, 12, 28])
        name = f"{title}_{stamp}.xlsx"
    elif kind == "user360":
        rows = user360.export_rows(f, keyword=keyword)
        ws.title = "用户360明细"
        _write(ws, ["用户协议ID", "城市", "电池产品", "所属线路", "用户ID", "姓名", "用户签约手机号",
                    "用户当前手机号", "协议类型", "协议状态",
                    "生效时间", "生效天数", "租期到期", "套餐名称", "押金金额(元)", "押金状态",
                    "押金创建时间", "押金到期时间", "押金预警", "划扣状态", "押金划扣失败原因",
                    "建议划扣金额(元)",
                    "近30天换电", "近90天换电", "月均频次", "近30天消费(元)", "近90天消费(元)", "累计消费(元)",
                    "电池SN", "电池状态", "电池在线", "电量(%)", "最后定位", "最后换电网点", "是否已归还柜机",
                    "最近接待备注", "接待时间", "用户标签", "评价得分", "评价次数", "综合处置建议"],
               [[r.get("agreement_id"), r.get("city"), r.get("product"), r.get("line"), r.get("user_id"),
                 r.get("user_name"), r.get("user_phone"), r.get("current_phone"),
                 r.get("agreement_type"), r.get("agreement_status"),
                 r.get("activation_time"), r.get("active_days"), r.get("rent_expire_time"),
                 r.get("package_name"), r.get("deposit_fee_yuan"), r.get("deposit_status"),
                 r.get("deposit_bind_str"), r.get("deposit_expire_str"),
                 r.get("deposit_warn") or r.get("deposit_warn_text"),
                 r.get("deduct_status_cn"), r.get("failure_name"), r.get("suggest_fee_yuan"),
                 r.get("n30"), r.get("n90"), r.get("monthly_freq"),
                 r.get("fee30_yuan"), r.get("fee90_yuan"), r.get("fee180_yuan"),
                 r.get("battery_sn"), r.get("battery_status"), r.get("battery_online"), r.get("battery_power"),
                 r.get("battery_location"), r.get("last_site"), "是" if r.get("returned") else "否",
                 r.get("note"), r.get("note_time"), "、".join(r.get("tags") or []),
                 r.get("score"), r.get("score_cnt"), r.get("advice")] for r in rows],
               [12, 14, 16, 18, 12, 12, 16, 16, 10, 10, 18, 10, 18, 18, 12, 12, 20, 20, 26, 12, 18, 16,
                10, 10, 10, 14, 14, 14,
                20, 10, 10, 8, 34, 18, 14, 40, 18, 30, 10, 10, 14])
        name = f"用户360明细_{stamp}.xlsx"
    else:
        raise ValueError(f"未知导出类型：{kind}")

    path = _out_path(name)
    wb.save(path)
    return path
